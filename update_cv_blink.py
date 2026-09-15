from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph

source = "/Users/alastairmcneill/Documents/GitHub/AlastairMcNeill.com/Alastair_McNeill_Yila_Project_Production_CV_updated.docx"
output = "/Users/alastairmcneill/Documents/GitHub/AlastairMcNeill.com/Alastair_McNeill_Creative_CV_2026.docx"

document = Document(source)

profile = next(p for p in document.paragraphs if p.text.startswith("British sound artist"))
profile.text = profile.text.replace("live musical direction and commercial composition", "live musical direction, theatre composition and commercial composition")

anchor = next(p for p in document.paragraphs if p.text.startswith("Selected collaborators and credits include"))
new_xml = deepcopy(anchor._p)
anchor._p.addnext(new_xml)
blink = Paragraph(new_xml, anchor._parent)
blink.text = "Current theatre composition work includes Blink."
blink.style = anchor.style

# The source CV contains a manual page break after the Garance production entry.
# Clear that specific break so the updated CV does not waste a page.
paragraphs = list(document.paragraphs)
for index, paragraph in enumerate(paragraphs):
    previous = paragraphs[index - 1].text if index else ""
    if (not paragraph.text.strip() and 'w:type="page"' in paragraph._p.xml
            and previous.startswith('Produced the full-length album')):
        paragraph._element.getparent().remove(paragraph._element)
        break

document.save(output)
print(output)
